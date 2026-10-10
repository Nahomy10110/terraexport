"""
Convierte la hoja de cálculo verificada en el archivo de datos del prototipo.

Uso (desde la carpeta raíz del repositorio):
    pip install openpyxl
    python scripts/excel_a_datos.py

Lee:      hoja-calculo/TerraExport_Modelo_Calculo.xlsx
Escribe:  datos/muestra.json (lo usan packages/motor y apps/web)

Importante: guarde el archivo en Excel (o LibreOffice) antes de ejecutar este
script. openpyxl lee los valores que Excel dejó calculados; si una celda con
fórmula sale vacía es porque el archivo no se guardó después de editarlo.
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
XLSX = RAIZ / "hoja-calculo" / "TerraExport_Modelo_Calculo.xlsx"
SALIDA = RAIZ / "datos" / "muestra.json"


def filas(ws, desde=5):
    """Devuelve las filas con datos (se detiene en la primera fila sin columna A)."""
    for fila in ws.iter_rows(min_row=desde, values_only=True):
        if fila[0] in (None, ""):
            break
        yield fila


def txt(v):
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d")
    return v


def main():
    if not XLSX.exists():
        sys.exit(f"No se encontró {XLSX}")
    wb = load_workbook(XLSX, data_only=True)

    p = wb["Parametros"]
    parametros = {
        "capitalInicial": p["B5"].value,
        "elasticidad": p["B6"].value,
        "rondas": p["B7"].value,
        "duracionMin": p["B8"].value,
        "origen": p["B9"].value,
    }

    fuentes = {f[0]: {"nombre": f[1], "url": f[2] or "", "uso": f[3] if len(f) > 3 and f[3] else ""} for f in filas(wb["Fuentes"])}

    paises = []
    for f in filas(wb["Paises"]):
        paises.append({
            "iso": f[0], "nombre": f[1], "moneda": f[2], "codigoMoneda": f[3],
            "unidadesPorUsd": f[4], "fuenteTc": f[5], "fechaCorte": txt(f[6]),
            "hofstede": {"distanciaPoder": f[7], "individualismo": f[8], "masculinidad": f[9],
                          "aversionIncertidumbre": f[10], "largoPlazo": f[11], "indulgencia": f[12]},
            "fuenteHofstede": f[13], "estado": f[14],
        })

    productos = [{"id": f[0], "nombre": f[1], "codigoSa": f[2], "unidad": f[3],
                  "pesoKg": f[4], "costoFob": f[5], "fuente": f[6], "estado": f[7], "nota": f[8] or ""}
                 for f in filas(wb["Productos"])]

    rutas = [{"iso": f[1], "medio": f[3], "fletePorKg": f[4], "pctSeguro": f[5],
              "diasTransito": f[6], "fuente": f[7], "estado": f[8], "nota": f[9] or ""}
             for f in filas(wb["Rutas"])]

    aranceles = [{"iso": f[1], "producto": f[2], "pct": f[5], "fuente": f[6],
                  "fechaCorte": txt(f[7]), "estado": f[8], "nota": f[9] or ""}
                 for f in filas(wb["Aranceles"])]

    demanda = [{"iso": f[1], "producto": f[2], "precioReferencia": f[5],
                "demandaBase": f[6], "fuente": f[7], "estado": f[8]}
               for f in filas(wb["Demanda"])]

    casos = []
    for f in filas(wb["Casos"]):
        if not str(f[0]).startswith("C"):
            break
        if f[17] is None:
            sys.exit("Las fórmulas de la hoja Casos no tienen valores calculados. "
                     "Abra el archivo en Excel, guárdelo y vuelva a ejecutar el script.")
        casos.append({
            "caso": f[0],
            "decision": {"iso": f[1], "producto": f[2], "medio": f[3], "volumen": f[4],
                          "precio": f[5], "equipos": f[6], "variacion": f[7], "capital": f[8]},
            "esperado": {"valorFob": f[17], "flete": f[18], "seguro": f[19], "valorCif": f[20],
                          "arancel": f[21], "costoTotal": f[22], "tcAplicado": f[23],
                          "costoLocal": f[24], "indicePrecio": f[25], "unidadesVendidas": f[26],
                          "ventas": f[27], "utilidad": f[28], "capitalFinal": f[29],
                          "valida": f[30], "explicacion": f[31]},
        })

    datos = {
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "parametros": parametros, "fuentes": fuentes, "paises": paises,
        "productos": productos, "rutas": rutas, "aranceles": aranceles,
        "demanda": demanda, "casos": casos,
    }
    SALIDA.parent.mkdir(exist_ok=True)
    SALIDA.write_text(json.dumps(datos, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Listo: {SALIDA.relative_to(RAIZ)} ({len(paises)} países, {len(productos)} productos, {len(casos)} casos)")


if __name__ == "__main__":
    main()
